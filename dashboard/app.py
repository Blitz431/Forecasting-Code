"""AutoStockAnalyzer — Streamlit Dashboard entry point (Settings landing page).

Run with:
    streamlit run dashboard/app.py
"""

from __future__ import annotations

import io
import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import streamlit as st

st.set_page_config(page_title="Settings", page_icon="🔐", layout="wide")

from dashboard.components.market_clock import render_market_clock
render_market_clock()
from dashboard.components.pipeline_progress import render_pipeline_progress
render_pipeline_progress()

st.title("🔐 Settings — API Keys & Secrets")
st.caption(
    "Values are written to the project-root `.env` file. The file is "
    "git-ignored so nothing you enter here is ever committed."
)
st.divider()

ROOT_DIR = Path(__file__).parent.parent
ENV_FILE = ROOT_DIR / ".env"
ENV_EXAMPLE = ROOT_DIR / ".env.example"

# Each entry: (env_var, display_label, sensitive?, help_text)
KEY_SPEC: list[tuple[str, str, bool, str]] = [
    ("ALPACA_API_KEY",       "Alpaca API Key",           True,  "From https://app.alpaca.markets/"),
    ("ALPACA_SECRET_KEY",    "Alpaca Secret Key",        True,  "Paired with the API key above."),
    ("ALPACA_BASE_URL",      "Alpaca Base URL",          False, "paper-api.alpaca.markets for paper trading."),
    ("FRED_API_KEY",         "FRED API Key",             True,  "From https://fred.stlouisfed.org/docs/api/api_key.html"),
    ("TRADING_MODE",         "Trading Mode",             False, "paper or live — NEVER set to live casually."),
    ("DISCORD_WEBHOOK_URL",  "Discord Webhook URL",      True,  "Optional — leave blank to disable Discord alerts."),
    ("SMTP_HOST",            "SMTP Host",                False, "e.g. smtp.gmail.com"),
    ("SMTP_PORT",            "SMTP Port",                False, "Usually 587."),
    ("SMTP_USER",            "SMTP Username",            False, "Full email address."),
    ("SMTP_PASSWORD",        "SMTP Password",            True,  "App-password recommended."),
    ("ALERT_EMAIL_TO",       "Alert Recipient Email",    False, "Where daily alerts get sent."),
    ("AUTO_RANK_ENABLED",    "Auto-run daily ranking",   False, "true or false. Runs the full pipeline automatically on the first dashboard open each day. Buttons still work."),
]


def _parse_env_file(path: Path) -> dict[str, str]:
    if not path.exists():
        return {}
    values: dict[str, str] = {}
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, val = line.partition("=")
        values[key.strip()] = val.strip()
    return values


def _write_env_file(path: Path, values: dict[str, str]) -> None:
    """Atomically write KEY=VALUE lines, preserving comment/ordering cues from .env.example."""
    lines: list[str] = []
    template_path = path if path.exists() else ENV_EXAMPLE
    written: set[str] = set()

    if template_path.exists():
        for raw in template_path.read_text(encoding="utf-8").splitlines():
            stripped = raw.strip()
            if stripped and not stripped.startswith("#") and "=" in stripped:
                key = stripped.split("=", 1)[0].strip()
                if key in values:
                    lines.append(f"{key}={values[key]}")
                    written.add(key)
                    continue
            lines.append(raw)

    for key, val in values.items():
        if key not in written:
            lines.append(f"{key}={val}")

    tmp_path = path.with_suffix(path.suffix + ".tmp")
    tmp_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    os.replace(tmp_path, path)


def _status_docx(values: dict[str, str]) -> bytes:
    """Build a .docx listing every key + SET/NOT SET status. Never contains raw values."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("AutoStockAnalyzer — API Key Inventory", level=1)
    doc.add_paragraph(
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        "This document lists every API key / secret the application expects "
        "and whether it is currently configured. Raw values are never included."
    )

    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Environment Variable"
    hdr[1].text = "Description"
    hdr[2].text = "Status"

    for key, label, sensitive, help_text in KEY_SPEC:
        row = table.add_row().cells
        row[0].text = key
        row[1].text = label + ((" — " + help_text) if help_text else "")
        row[2].text = "SET" if values.get(key) else "NOT SET"

    doc.add_paragraph("")
    note = doc.add_paragraph()
    run = note.add_run(
        "Security note: only variable names and SET/NOT SET status are exported. "
        "No raw key values are included in this file."
    )
    run.font.size = Pt(9)
    run.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------#
# UI
# ---------------------------------------------------------------------------#

current = _parse_env_file(ENV_FILE)

if "settings_reveal" not in st.session_state:
    st.session_state.settings_reveal = False

col_a, col_b = st.columns([1, 3])
with col_a:
    st.session_state.settings_reveal = st.toggle(
        "Reveal values", value=st.session_state.settings_reveal,
        help="Local-only — values are never logged."
    )
with col_b:
    if ENV_FILE.exists():
        set_count = sum(1 for k, *_ in KEY_SPEC if current.get(k))
        st.info(f"`.env` found — {set_count}/{len(KEY_SPEC)} keys configured.")
    else:
        st.warning("`.env` does not exist yet — saving will create it.")

st.divider()

with st.form("env_form"):
    edits: dict[str, str] = {}
    for key, label, sensitive, help_text in KEY_SPEC:
        current_val = current.get(key, "")
        input_type = "default" if (not sensitive or st.session_state.settings_reveal) else "password"
        edits[key] = st.text_input(
            label,
            value=current_val,
            type=input_type,
            help=f"Env var: `{key}` — {help_text}",
            key=f"env_input_{key}",
        )

    submitted = st.form_submit_button("💾 Save to .env", type="primary")
    if submitted:
        merged = {**current, **{k: v for k, v in edits.items() if v is not None}}
        try:
            _write_env_file(ENV_FILE, merged)
            st.success(f"Saved {len(edits)} entries to {ENV_FILE}")
            st.info("Reload the dashboard (`R`) for new values to take effect.")
        except Exception as exc:
            st.error(f"Write failed: {exc}")

st.divider()

st.subheader("Export Key Inventory")
st.write(
    "Generate a Word document listing every expected key and whether it's "
    "configured. **Raw values are never included** — the file is safe to share."
)

if st.button("📄 Build Word inventory"):
    try:
        latest = _parse_env_file(ENV_FILE)
        docx_bytes = _status_docx(latest)
        st.download_button(
            label="⬇ Download api_keys_inventory.docx",
            data=docx_bytes,
            file_name=f"api_keys_inventory_{datetime.now().strftime('%Y%m%d')}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except ImportError:
        st.error("Missing dependency: install `python-docx` (`pip install python-docx`).")
    except Exception as exc:
        st.error(f"Word export failed: {exc}")

st.divider()

# ---------------------------------------------------------------------------#
# Data Pipeline (scraper)
# ---------------------------------------------------------------------------#

st.subheader("Data Pipeline — Scraper")
st.write(
    "Download or update stock prices (S&P 500), FRED macro data, and dividends. "
    "**Full backfill takes 30–60 minutes** for all 500 tickers from 2015. "
    "Run it once, then use Incremental Update daily."
)

col_s1, col_s2 = st.columns(2)
with col_s1:
    custom_tickers_scrape = st.text_input(
        "Custom tickers (leave blank for full S&P 500)",
        placeholder="AAPL,MSFT,NVDA",
        key="scrape_tickers",
    )
with col_s2:
    scrape_prices_only = st.checkbox("Prices only (skip FRED macro + dividends)", key="scrape_prices_only")

btn_col1, btn_col2, btn_col3 = st.columns(3)

with btn_col1:
    run_incremental = st.button("🔄 Incremental Update", help="Download only missing/new data")
with btn_col2:
    run_backfill = st.button("⬇ Full Backfill (2015→now)", type="primary",
                              help="Wipe and re-download everything — slow")
with btn_col3:
    run_macro = st.button("📊 FRED Macro Only", help="Refresh GDP, CPI, VIX, etc. from FRED")

def _run_scraper(extra_args: list[str]) -> None:
    cli_path = ROOT_DIR / "cli" / "scrape.py"
    cmd = [sys.executable, str(cli_path)] + extra_args
    if custom_tickers_scrape.strip():
        from src.utils.input_sanitize import clean_ticker_list
        tickers, rejected = clean_ticker_list(custom_tickers_scrape)
        if rejected:
            st.warning(f"Ignored invalid tickers: {', '.join(rejected)}")
        if tickers:
            cmd += ["--tickers", ",".join(tickers)]
    if scrape_prices_only:
        cmd += ["--prices-only"]

    log_box = st.empty()
    lines: list[str] = []
    try:
        proc = subprocess.Popen(
            cmd,
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            cwd=str(ROOT_DIR),
        )
        for line in proc.stdout:
            lines.append(line.rstrip())
            log_box.text_area("Output", "\n".join(lines[-80:]), height=320)
        proc.wait()
        if proc.returncode == 0:
            st.success("Scrape finished successfully.")
        else:
            st.error(f"Scraper exited with code {proc.returncode}.")
    except Exception as exc:
        st.error(f"Scraper failed to start: {exc}")

if run_incremental:
    with st.spinner("Running incremental update..."):
        _run_scraper([])

if run_backfill:
    with st.spinner("Running full backfill — this will take 30-60 minutes..."):
        _run_scraper(["--backfill"])

if run_macro:
    with st.spinner("Fetching FRED macro data..."):
        _run_scraper(["--macro-only"])

st.divider()

# ---------------------------------------------------------------------------#
# Analysis Pipeline (ranker + ML)
# ---------------------------------------------------------------------------#

st.subheader("Analysis Pipeline — Ranker & ML")
st.write(
    "**Phase 1** fast-ranks all tickers using indicators, forecasts, and sentiment. "
    "**Phase 2** trains the full ML models on the top candidates, then re-ranks them. "
    "Run after a scrape update to refresh the trading signals."
)
st.caption(
    "The full pipeline runs automatically in the background once per day on the first "
    "dashboard open (set AUTO_RANK_ENABLED=false above to turn this off). If it fails, "
    "a fast rank runs instead."
)

from src.pipeline_runner.launcher import (
    DEFAULT_POOL,
    DEFAULT_TOP,
    reap_children,
    start_pipeline_run,
)
from src.pipeline_runner.run_state import (
    get_run_paths,
    is_run_active,
    read_log_tail,
    read_status,
)
from src.utils.logging import setup_logger

logger = setup_logger(__name__)

reap_children()
try:
    _run_paths = get_run_paths()
    pipeline_active = is_run_active(_run_paths)
except Exception as _exc:
    # Unreadable/read-only data dir: treat as idle, launcher reports spawn errors.
    logger.warning(f"Pipeline run-state lookup failed: {_exc}")
    _run_paths = None
    pipeline_active = False

pipe_col1, pipe_col2 = st.columns(2)
with pipe_col1:
    pipe_pool = st.number_input(
        "Phase-1 pool size (fed into ML training)",
        min_value=10, max_value=503, value=DEFAULT_POOL, step=10,
        help="Top-N fast-ranked tickers that get the full ML treatment.",
    )
with pipe_col2:
    pipe_top = st.number_input(
        "Final top-N to show",
        min_value=5, max_value=200, value=DEFAULT_TOP, step=5,
        help="How many ranked results to display and cache.",
    )

pipe_btn_col1, pipe_btn_col2 = st.columns(2)
with pipe_btn_col1:
    run_phase1_only = st.button(
        "⚡ Fast Rank Only (Phase 1)",
        disabled=pipeline_active,
        help="Quick no-ML ranking (~30s). Replaces the current rankings, "
             "including today's ML rankings if any.",
    )
with pipe_btn_col2:
    run_full_pipeline = st.button(
        "🧠 Full Analysis Pipeline (Phase 1 + 2)",
        type="primary",
        disabled=pipeline_active,
        help="Phase 1 + ML training. If a fast rank already ran today, only the ML part runs.",
    )


def _start_pipeline(mode: str) -> None:
    res = start_pipeline_run(mode, int(pipe_pool), int(pipe_top), trigger="button")
    if res.started:
        st.success(res.message)
        st.rerun()
    else:
        st.warning(res.message)


if run_phase1_only:
    _start_pipeline("phase1")

if run_full_pipeline:
    _start_pipeline("full")


def _fmt_local(value: object) -> str:
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if dt.tzinfo is None:
            from datetime import timezone
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.astimezone().strftime("%Y-%m-%d %H:%M:%S")
    except (ValueError, TypeError):
        return "-"


_last = read_status(_run_paths) if _run_paths is not None else None
if _last:
    _parts = [
        f"state: {_last.get('state', '?')}",
        f"mode: {_last.get('mode', '?')}",
        f"trigger: {_last.get('trigger', '?')}",
        f"phase 1 reused: {'yes' if _last.get('phase1_reused') else 'no'}",
        f"started: {_fmt_local(_last.get('started_at'))}",
    ]
    if _last.get("finished_at"):
        _parts.append(f"finished: {_fmt_local(_last.get('finished_at'))}")
    if _last.get("state") == "failed" and _last.get("error"):
        _parts.append(f"error: {str(_last['error']).splitlines()[0][:200]}")
    st.caption("Last run - " + " | ".join(_parts))
    if _last.get("run_id"):
        with st.expander("Last run log"):
            st.code("\n".join(read_log_tail(_run_paths, str(_last["run_id"]))) or "(no log output)")

st.divider()

# ---------------------------------------------------------------------------#
# Shut Down
# ---------------------------------------------------------------------------#

st.subheader("Shut Down")
st.write(
    "Stops the Streamlit server and closes the terminal (command prompt) it was "
    "launched from. Save any open work before clicking."
)

if "shutdown_confirm" not in st.session_state:
    st.session_state.shutdown_confirm = False

if not st.session_state.shutdown_confirm:
    if st.button("🔴 Shut Down Server", type="secondary"):
        st.session_state.shutdown_confirm = True
        st.rerun()
else:
    st.warning("Are you sure? This will stop Streamlit and close the terminal.")
    c1, c2 = st.columns(2)
    with c1:
        if st.button("Yes, shut down", type="primary"):
            st.info("Shutting down...")
            os._exit(0)
    with c2:
        if st.button("Cancel"):
            st.session_state.shutdown_confirm = False
            st.rerun()
